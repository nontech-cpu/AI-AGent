import io
import os
import re
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def build_contacts_excel(contacts, tasks=None) -> bytes:
    """Builds a formatted multi-sheet Excel spreadsheet containing contacts and tasks."""
    wb = openpyxl.Workbook()
    # Sheet 1: Contacts & Networking
    ws = wb.active
    ws.title = "Contacts & Networking"

    headers = [
        "Full Name", "Company", "Job Title / Role", "Email", "Phone",
        "Met At / Location", "Context & Notes", "Commitments / Action Items", "Saved Date"
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for c in contacts:
        ws.append([
            c.get("name") or "",
            c.get("company") or "",
            c.get("role_title") or "",
            c.get("email") or "",
            c.get("phone") or "",
            c.get("met_at") or "",
            c.get("context_notes") or "",
            c.get("commitments") or "",
            str(c.get("created_at") or "")[:10]
        ])

    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 2), min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 14)

    # Sheet 2: Tasks & Reminders (if any)
    if tasks:
        ws_tasks = wb.create_sheet(title="Tasks & Reminders")
        task_headers = ["Category", "Title", "Amount", "Due Date", "Time"]
        ws_tasks.append(task_headers)
        for col_idx in range(1, len(task_headers) + 1):
            cell = ws_tasks.cell(row=1, column=col_idx)
            cell.fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for t in tasks:
            ws_tasks.append([
                t[0] if len(t) > 0 else "",
                t[1] if len(t) > 1 else "",
                t[2] if len(t) > 2 else "",
                str(t[3]) if len(t) > 3 and t[3] else "",
                str(t[4])[:5] if len(t) > 4 and t[4] else ""
            ])
        for col in ws_tasks.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_tasks.column_dimensions[col_letter].width = max(min(max_len + 4, 40), 14)

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def build_styled_pdf(title: str, sections: list[dict], author: str = "PeculiarAI") -> bytes:
    """
    Builds an executive PDF document with title, timestamp, and styled sections.
    sections is a list of dicts: {"heading": str, "paragraphs": [str], "table": list[list[str]] (optional)}
    """
    bio = io.BytesIO()
    doc = SimpleDocTemplate(
        bio,
        pagesize=letter,
        leftMargin=45,
        rightMargin=45,
        topMargin=45,
        bottomMargin=45
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1F4E79'),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#595959'),
        spaceAfter=15,
    )
    h2_style = ParagraphStyle(
        'Heading2Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#2E75B6'),
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        'BodyCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#262626'),
        spaceAfter=8,
    )

    story = [
        Paragraph(title, title_style),
        Paragraph(f"Prepared by {author} • Generated on {datetime.now().strftime('%d %B %Y, %I:%M %p')}", subtitle_style),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1F4E79'), spaceBefore=2, spaceAfter=14),
    ]

    for sec in sections:
        if sec.get("heading"):
            story.append(Paragraph(sec["heading"], h2_style))
        for p in sec.get("paragraphs", []):
            story.append(Paragraph(p, body_style))
        
        # Optional table in section
        if sec.get("table"):
            t_data = []
            for row in sec["table"]:
                t_row = [Paragraph(str(cell), body_style) for cell in row]
                t_data.append(t_row)
            if t_data:
                table = Table(t_data, colWidths=[120] + [380 // (len(t_data[0]) - 1)] * (len(t_data[0]) - 1) if len(t_data[0]) > 1 else [500])
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F2F4F7')),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D9D9D9')),
                ]))
                story.append(table)
                story.append(Spacer(1, 10))

    doc.build(story)
    return bio.getvalue()
